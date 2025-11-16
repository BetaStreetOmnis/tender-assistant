#!/usr/bin/env python3
"""
AI标书助理系统 - 完整集成版
集成功能：
1. DocuGen模板识别和填充
2. AI智能大纲和内容生成
3. 简化的RAG知识库功能
4. 智谱GLM-4真实AI能力
"""
import uvicorn
from fastapi import FastAPI, HTTPException, Form, status, UploadFile, File
from fastapi.responses import FileResponse
from pydantic import BaseModel
from typing import Optional, List, Dict, Any
import json
import time
from datetime import datetime
import requests
import re
import os
import uuid
from pathlib import Path

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    print("⚠️ python-docx未安装，模板功能将不可用")

app = FastAPI(
    title="AI标书助理系统 - 完整集成版",
    description="集成DocuGen模板、AI生成、RAG知识库功能",
    version="4.0.0"
)

# 添加CORS中间件
from fastapi.middleware.cors import CORSMiddleware
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# 智谱GLM配置
ZHIPU_CONFIG = {
    'api_key': 'ba4c4317831947cbb697e7ad68393308.qdUFMMBgGg8wtqHc',
    'base_url': 'https://open.bigmodel.cn/api/paas/v4/',
    'model': 'glm-4'
}

# 创建必要的目录
TEMPLATES_DIR = Path("templates")
DOWNLOADS_DIR = Path("downloads")
KNOWLEDGE_DIR = Path("knowledge")  # 知识库文件存储
TEMPLATES_DIR.mkdir(exist_ok=True)
DOWNLOADS_DIR.mkdir(exist_ok=True)
KNOWLEDGE_DIR.mkdir(exist_ok=True)

# ============ 数据模型 ============

class AnalyzeRequest(BaseModel):
    tender_content: str = ''

class GenerateFromTemplateRequest(BaseModel):
    template_id: str
    variables: Dict[str, Any]

class GenerateOutlineRequest(BaseModel):
    topic: str
    key_points: Optional[str] = None
    use_knowledge: bool = False  # 是否使用知识库增强

class GenerateSectionRequest(BaseModel):
    outline_section: str
    context: Optional[str] = None
    use_knowledge: bool = False

class KnowledgeUploadRequest(BaseModel):
    title: str
    content: str
    category: Optional[str] = "通用"

# ============ 简化的知识库系统 ============

class SimpleKnowledgeBase:
    """简化的知识库系统（基于文件存储）"""

    def __init__(self, storage_dir: Path):
        self.storage_dir = storage_dir
        self.index_file = storage_dir / "index.json"
        self._load_index()

    def _load_index(self):
        """加载知识库索引"""
        if self.index_file.exists():
            with open(self.index_file, 'r', encoding='utf-8') as f:
                self.index = json.load(f)
        else:
            self.index = {"documents": []}
            self._save_index()

    def _save_index(self):
        """保存知识库索引"""
        with open(self.index_file, 'w', encoding='utf-8') as f:
            json.dump(self.index, f, ensure_ascii=False, indent=2)

    def add_document(self, title: str, content: str, category: str = "通用") -> str:
        """添加文档到知识库"""
        doc_id = str(uuid.uuid4())[:8]
        doc_file = self.storage_dir / f"{doc_id}.txt"

        # 保存文档内容
        with open(doc_file, 'w', encoding='utf-8') as f:
            f.write(content)

        # 更新索引
        self.index["documents"].append({
            "id": doc_id,
            "title": title,
            "category": category,
            "file": str(doc_file),
            "created_at": datetime.now().isoformat(),
            "size": len(content)
        })
        self._save_index()

        return doc_id

    def search(self, query: str, top_k: int = 3) -> List[Dict[str, Any]]:
        """简单的关键词搜索"""
        results = []
        query_lower = query.lower()

        for doc in self.index["documents"]:
            doc_file = Path(doc["file"])
            if doc_file.exists():
                with open(doc_file, 'r', encoding='utf-8') as f:
                    content = f.read()

                # 简单的关键词匹配得分
                score = 0
                for word in query_lower.split():
                    if word in content.lower():
                        score += content.lower().count(word)

                if score > 0:
                    results.append({
                        "id": doc["id"],
                        "title": doc["title"],
                        "category": doc["category"],
                        "content": content[:500],  # 返回前500字符
                        "score": score
                    })

        # 按得分排序并返回top_k
        results.sort(key=lambda x: x["score"], reverse=True)
        return results[:top_k]

    def get_all_documents(self) -> List[Dict[str, Any]]:
        """获取所有文档"""
        return self.index["documents"]

    def delete_document(self, doc_id: str) -> bool:
        """删除文档"""
        for i, doc in enumerate(self.index["documents"]):
            if doc["id"] == doc_id:
                # 删除文件
                doc_file = Path(doc["file"])
                if doc_file.exists():
                    doc_file.unlink()

                # 从索引中移除
                self.index["documents"].pop(i)
                self._save_index()
                return True
        return False

# 初始化知识库
knowledge_base = SimpleKnowledgeBase(KNOWLEDGE_DIR)

# ============ AI调用函数 ============

def call_zhipu_ai(prompt: str, system_prompt: str = None, max_tokens: int = 2000) -> dict:
    """调用智谱GLM-4进行分析"""
    try:
        url = f"{ZHIPU_CONFIG['base_url']}chat/completions"
        headers = {
            'Authorization': f"Bearer {ZHIPU_CONFIG['api_key']}",
            'Content-Type': 'application/json'
        }

        messages = []
        if system_prompt:
            messages.append({
                'role': 'system',
                'content': system_prompt
            })
        messages.append({
            'role': 'user',
            'content': prompt
        })

        payload = {
            'model': ZHIPU_CONFIG['model'],
            'messages': messages,
            'temperature': 0.7,
            'max_tokens': max_tokens
        }

        response = requests.post(url, headers=headers, json=payload, timeout=60)
        response.raise_for_status()

        result = response.json()

        # 提取AI响应内容
        if 'choices' in result and len(result['choices']) > 0:
            choice = result['choices'][0]
            message = choice.get('message', {})
            content = message.get('content', '')

            if not content and 'reasoning_content' in message:
                content = message['reasoning_content']

            return {
                'success': True,
                'content': content,
                'model': result.get('model', ZHIPU_CONFIG['model']),
                'usage': result.get('usage', {})
            }
    except requests.exceptions.Timeout:
        return {'success': False, 'content': '', 'error': 'API请求超时'}
    except Exception as e:
        return {'success': False, 'content': '', 'error': f'API调用失败: {str(e)}'}

    return {'success': False, 'content': '', 'error': '未知错误'}

# ============ DocuGen模板功能 ============

def extract_variables_from_template(template_path: str) -> List[Dict[str, str]]:
    """从Word模板中提取所有变量"""
    if not DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="模板功能不可用：缺少python-docx库")

    if not os.path.exists(template_path):
        raise ValueError(f"模板文件不存在: {template_path}")

    try:
        doc = Document(template_path)
        variables = []
        found_vars = set()
        var_pattern = re.compile(r'【([^】]+?)】|\{\{([^}]+?)\}\}')

        for paragraph in doc.paragraphs:
            matches = var_pattern.findall(paragraph.text)
            for match_tuple in matches:
                var_name = next((v for v in match_tuple if v), None)
                if var_name:
                    var_name = var_name.strip()
                    if var_name and var_name not in found_vars:
                        found_vars.add(var_name)
                        variables.append({
                            "name": var_name,
                            "label": var_name,
                            "type": "text",
                            "required": True,
                            "description": f"{var_name}的值"
                        })

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        matches = var_pattern.findall(paragraph.text)
                        for match_tuple in matches:
                            var_name = next((v for v in match_tuple if v), None)
                            if var_name:
                                var_name = var_name.strip()
                                if var_name and var_name not in found_vars:
                                    found_vars.add(var_name)
                                    variables.append({
                                        "name": var_name,
                                        "label": var_name,
                                        "type": "text",
                                        "required": True,
                                        "description": f"{var_name}的值"
                                    })

        return variables
    except Exception as e:
        raise Exception(f"提取模板变量失败: {str(e)}")

def fill_template(template_path: str, data: Dict[str, Any]) -> str:
    """填充模板并生成新文档"""
    if not DOCX_AVAILABLE:
        raise HTTPException(status_code=500, detail="模板功能不可用：缺少python-docx库")

    try:
        doc = Document(template_path)

        for paragraph in doc.paragraphs:
            _replace_text_in_paragraph(paragraph, data)

        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        _replace_text_in_paragraph(paragraph, data)

        filename = f"generated_{uuid.uuid4().hex[:8]}.docx"
        file_path = DOWNLOADS_DIR / filename
        doc.save(str(file_path))

        return str(file_path)
    except Exception as e:
        raise Exception(f"填充模板失败: {str(e)}")

def _replace_text_in_paragraph(paragraph, data: Dict[str, Any]) -> None:
    """在段落中替换占位符"""
    for key, value in data.items():
        placeholders = [f"【{key}】", f"{{{{{key}}}}}"]
        for p_text in placeholders:
            while p_text in paragraph.text:
                runs = paragraph.runs
                full_text = "".join(r.text for r in runs)

                if p_text not in full_text:
                    break

                start_idx = full_text.find(p_text)
                end_idx = start_idx + len(p_text)

                run_cursor = 0
                start_run, end_run = None, None
                start_run_idx, end_run_idx = -1, -1
                start_offset, end_offset = -1, -1

                for i, run in enumerate(runs):
                    run_len = len(run.text)
                    if start_run is None and run_cursor + run_len > start_idx:
                        start_run = run
                        start_run_idx = i
                        start_offset = start_idx - run_cursor
                    if end_run is None and run_cursor + run_len >= end_idx:
                        end_run = run
                        end_run_idx = i
                        end_offset = end_idx - run_cursor
                        break
                    run_cursor += run_len

                if start_run is None or end_run is None:
                    break

                if start_run_idx == end_run_idx:
                    start_run.text = start_run.text[:start_offset] + str(value) + end_run.text[end_offset:]
                else:
                    start_run.text = start_run.text[:start_offset] + str(value)
                    end_run.text = end_run.text[end_offset:]
                    for i in range(start_run_idx + 1, end_run_idx):
                        runs[i].text = ""

# ... (继续添加API端点)
# ============ API端点 ============

@app.get('/')
async def root():
    """根路径"""
    return {
        "message": "AI标书助理系统 - 完整集成版",
        "version": "4.0.0",
        "status": "running",
        "features": [
            "✅ DocuGen模板识别和填充",
            "✅ AI智能大纲生成",
            "✅ AI智能内容生成",
            "✅ 简化RAG知识库",
            "✅ 智谱GLM-4真实AI",
            "✅ 招标文档分析"
        ]
    }

@app.get('/api/v1/health')
async def health_check():
    """健康检查"""
    return {
        'code': 200,
        'message': 'success',
        'data': {
            'status': 'healthy',
            'service': 'AI标书助理系统 - 完整集成版',
            'timestamp': datetime.now().isoformat(),
            'version': '4.0.0',
            'knowledge_docs': len(knowledge_base.get_all_documents())
        }
    }

# ============ 招标文档分析 ============

@app.post('/api/v1/bidding/analyze')
async def analyze(request: Optional[AnalyzeRequest] = None, tender_content: Optional[str] = Form(None)):
    """分析招标文档"""
    try:
        content = tender_content if tender_content is not None else (request.tender_content if request else None)

        if not content or not content.strip():
            raise HTTPException(
                status_code=status.HTTP_400_BAD_REQUEST,
                detail='缺少招标文档内容'
            )

        # 使用AI分析
        prompt = f"""请分析以下招标文档，提供详细的投标建议：

{content}

请从以下几个方面进行分析：
1. 项目概述和关键特点
2. 技术要求分析
3. 商务条款要点
4. 资质和业绩要求
5. 投标策略建议
6. 潜在风险和应对措施

请用清晰的结构化格式输出，每个要点用项目符号列出。"""

        ai_result = call_zhipu_ai(
            prompt,
            system_prompt='你是一位专业的招投标分析专家，擅长分析招标文档并提供专业的投标建议。请用中文回答。'
        )

        if ai_result['success']:
            analysis_text = ai_result['content']

            # 提取关键要点
            key_points = []
            for line in analysis_text.split('\n'):
                line = line.strip()
                if line and (line.startswith('•') or line.startswith('-') or line.startswith('*') or
                            any(line.startswith(f'{i}.') for i in range(1, 10))):
                    point = line.lstrip('•-*').strip()
                    if point and len(point) > 5:
                        key_points.append(point)

            if not key_points:
                key_points = [
                    '准备详细的技术实施方案',
                    '展示相关专业人员资质和经验',
                    '制定科学的项目管理计划',
                    '建立完善的质量控制体系',
                    '提供全面的售后技术支持'
                ]

            return {
                'code': 200,
                'message': 'success',
                'data': {
                    'id': str(int(time.time())),
                    'analysis': analysis_text,
                    'key_points': key_points[:8],
                    'created_at': datetime.now().isoformat(),
                    'model_used': f"智谱{ai_result.get('model', 'GLM-4')}",
                    'tokens_used': ai_result.get('usage', {}).get('total_tokens', 0),
                    'ai_confidence': 'high',
                    'analysis_type': 'ai_powered_analysis'
                }
            }
        else:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"AI分析失败: {ai_result.get('error', '未知错误')}"
            )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'分析时发生错误: {str(e)}'
        )

# ============ AI智能大纲生成 ============

@app.post('/api/v1/bidding/generate-outline')
async def generate_outline_endpoint(request: GenerateOutlineRequest):
    """生成项目大纲（支持RAG增强）"""
    try:
        # 构建基础prompt
        prompt = f"请为主题'{request.topic}'生成一个详细的、结构化的大纲。"

        # 如果有关键点
        if request.key_points:
            prompt += f"\n\n请确保大纲围绕以下关键点展开：\n{request.key_points}"

        # 如果启用知识库
        rag_context = ""
        if request.use_knowledge:
            kb_results = knowledge_base.search(request.topic, top_k=3)
            if kb_results:
                rag_context = "\n\n参考知识库内容：\n"
                for result in kb_results:
                    rag_context += f"\n【{result['title']}】\n{result['content']}\n"
                prompt += rag_context

        prompt += "\n\n请使用Markdown格式，用#、##、###表示不同层级的标题。"

        ai_result = call_zhipu_ai(
            prompt,
            system_prompt='你是一个专业的大纲生成助手，擅长创建结构化、有逻辑性的大纲。',
            max_tokens=3000
        )

        if ai_result['success']:
            return {
                'code': 200,
                'message': 'success',
                'data': {
                    'id': str(int(time.time())),
                    'outline': ai_result['content'],
                    'created_at': datetime.now().isoformat(),
                    'model_used': ai_result.get('model'),
                    'tokens_used': ai_result.get('usage', {}).get('total_tokens', 0),
                    'knowledge_enhanced': request.use_knowledge and len(rag_context) > 0
                }
            }
        else:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"大纲生成失败: {ai_result.get('error')}"
            )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'生成大纲时发生错误: {str(e)}'
        )

# ============ AI智能章节生成 ============

@app.post('/api/v1/bidding/generate-section')
async def generate_section_endpoint(request: GenerateSectionRequest):
    """根据大纲生成章节内容（支持RAG增强）"""
    try:
        prompt = f"请根据以下大纲部分，生成详细的文档内容：\n\n{request.outline_section}"

        if request.context:
            prompt += f"\n\n上下文信息：\n{request.context}"

        # 如果启用知识库
        if request.use_knowledge:
            kb_results = knowledge_base.search(request.outline_section, top_k=2)
            if kb_results:
                prompt += "\n\n参考知识库内容：\n"
                for result in kb_results:
                    prompt += f"\n【{result['title']}】\n{result['content']}\n"

        prompt += "\n\n请生成结构清晰、内容详实的文档内容。"

        ai_result = call_zhipu_ai(
            prompt,
            system_prompt='你是一个专业的内容写作助手，擅长根据大纲扩展生成高质量的文档内容。',
            max_tokens=3000
        )

        if ai_result['success']:
            return {
                'code': 200,
                'message': 'success',
                'data': {
                    'section_content': ai_result['content'],
                    'created_at': datetime.now().isoformat(),
                    'tokens_used': ai_result.get('usage', {}).get('total_tokens', 0)
                }
            }
        else:
            raise HTTPException(
                status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
                detail=f"章节生成失败: {ai_result.get('error')}"
            )

    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'生成章节时发生错误: {str(e)}'
        )

# ============ 知识库管理 ============

@app.post('/api/v1/knowledge/upload')
async def upload_knowledge(request: KnowledgeUploadRequest):
    """上传知识到知识库"""
    try:
        doc_id = knowledge_base.add_document(
            title=request.title,
            content=request.content,
            category=request.category or "通用"
        )

        return {
            'code': 200,
            'message': 'success',
            'data': {
                'doc_id': doc_id,
                'title': request.title,
                'created_at': datetime.now().isoformat()
            }
        }
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'上传知识失败: {str(e)}'
        )

@app.get('/api/v1/knowledge/list')
async def list_knowledge():
    """获取知识库列表"""
    try:
        docs = knowledge_base.get_all_documents()
        return {
            'code': 200,
            'message': 'success',
            'data': {
                'documents': docs,
                'total': len(docs)
            }
        }
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'获取知识库列表失败: {str(e)}'
        )

@app.get('/api/v1/knowledge/search')
async def search_knowledge(query: str, top_k: int = 3):
    """搜索知识库"""
    try:
        results = knowledge_base.search(query, top_k=top_k)
        return {
            'code': 200,
            'message': 'success',
            'data': {
                'results': results,
                'query': query,
                'count': len(results)
            }
        }
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'搜索知识库失败: {str(e)}'
        )

@app.delete('/api/v1/knowledge/{doc_id}')
async def delete_knowledge(doc_id: str):
    """删除知识库文档"""
    try:
        success = knowledge_base.delete_document(doc_id)
        if success:
            return {
                'code': 200,
                'message': 'success',
                'data': {'doc_id': doc_id}
            }
        else:
            raise HTTPException(status_code=404, detail='文档不存在')
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'删除知识库文档失败: {str(e)}'
        )

# ============ DocuGen模板管理 ============

@app.get('/api/v1/template/list')
async def get_template_list():
    """获取模板列表"""
    templates = []
    if TEMPLATES_DIR.exists():
        for idx, template_file in enumerate(TEMPLATES_DIR.glob("*.docx"), start=1):
            templates.append({
                'id': str(idx),
                'name': template_file.stem,
                'filename': template_file.name,
                'description': f'{template_file.stem}模板',
                'category': '通用',
                'created_at': datetime.fromtimestamp(template_file.stat().st_ctime).isoformat()
            })

    if not templates:
        templates = [
            {
                'id': '1',
                'name': '智能建筑投标模板',
                'filename': 'smart_building_template.docx',
                'description': '适用于智能建筑项目的专业投标模板',
                'category': '建筑工程',
                'created_at': '2024-01-01T00:00:00Z'
            }
        ]

    return {
        'code': 200,
        'message': 'success',
        'data': {'templates': templates}
    }

@app.get('/api/v1/template/variables')
async def get_template_variables(template_id: str = None):
    """获取模板变量"""
    try:
        if not template_id:
            return {
                'code': 200,
                'message': 'success',
                'data': {
                    'variables': [
                        {
                            'name': 'project_name',
                            'label': '项目名称',
                            'type': 'text',
                            'required': True,
                            'description': '招标项目的完整名称'
                        }
                    ]
                }
            }

        template_files = list(TEMPLATES_DIR.glob("*.docx"))
        if int(template_id) > len(template_files):
            raise HTTPException(status_code=404, detail="模板不存在")

        template_path = str(template_files[int(template_id) - 1])

        if DOCX_AVAILABLE:
            variables = extract_variables_from_template(template_path)
        else:
            variables = []

        return {
            'code': 200,
            'message': 'success',
            'data': {'variables': variables}
        }
    except Exception as e:
        return {
            'code': 200,
            'message': 'success',
            'data': {'variables': []}
        }

@app.post('/api/v1/template/generate-from-template')
async def generate_from_template_endpoint(request: GenerateFromTemplateRequest):
    """基于模板生成文档"""
    try:
        if not DOCX_AVAILABLE:
            raise HTTPException(status_code=500, detail="模板功能不可用：需要安装python-docx")

        template_files = list(TEMPLATES_DIR.glob("*.docx"))
        template_idx = int(request.template_id) - 1

        if template_idx < 0 or template_idx >= len(template_files):
            raise HTTPException(status_code=404, detail="模板不存在")

        template_path = str(template_files[template_idx])
        output_path = fill_template(template_path, request.variables)

        return {
            'code': 200,
            'message': 'success',
            'data': {
                'file_path': output_path,
                'filename': os.path.basename(output_path),
                'download_url': f'/api/v1/template/download/{os.path.basename(output_path)}'
            }
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'生成文档失败: {str(e)}'
        )

@app.get('/api/v1/template/download/{filename}')
async def download_generated_file(filename: str):
    """下载生成的文档"""
    file_path = DOWNLOADS_DIR / filename
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="文件不存在")

    return FileResponse(
        path=str(file_path),
        filename=filename,
        media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )

@app.post('/api/v1/template/upload')
async def upload_template(file: UploadFile = File(...)):
    """上传模板文件"""
    try:
        if not file.filename.endswith('.docx'):
            raise HTTPException(status_code=400, detail="只支持.docx格式")

        file_path = TEMPLATES_DIR / file.filename
        with open(file_path, 'wb') as f:
            content = await file.read()
            f.write(content)

        if DOCX_AVAILABLE:
            variables = extract_variables_from_template(str(file_path))
        else:
            variables = []

        return {
            'code': 200,
            'message': 'success',
            'data': {
                'template_name': file.filename,
                'variables': variables
            }
        }
    except HTTPException:
        raise
    except Exception as e:
        raise HTTPException(
            status_code=status.HTTP_500_INTERNAL_SERVER_ERROR,
            detail=f'上传模板失败: {str(e)}'
        )

# ============ 投标响应列表（前端兼容） ============

@app.get('/api/v1/bidding/responses')
async def get_responses(page: int = 1, page_size: int = 20):
    """获取投标响应列表"""
    return {
        'code': 200,
        'message': 'success',
        'data': {
            'items': [],
            'total': 0,
            'page': page,
            'page_size': page_size
        }
    }

@app.get('/api/v1/')
async def api_root():
    """API根路径"""
    return {
        'code': 200,
        'message': 'success',
        'data': {
            'service': 'AI标书助理系统API',
            'version': '4.0.0',
            'endpoints': [
                '/api/v1/bidding/analyze',
                '/api/v1/bidding/generate-outline',
                '/api/v1/bidding/generate-section',
                '/api/v1/bidding/responses',
                '/api/v1/knowledge/upload',
                '/api/v1/knowledge/list',
                '/api/v1/knowledge/search',
                '/api/v1/template/list',
                '/api/v1/template/variables',
                '/api/v1/template/generate-from-template',
                '/api/v1/template/upload'
            ]
        }
    }

# ============ 启动服务 ============

if __name__ == "__main__":
    print("🚀 启动AI标书助理系统 - 完整集成版 v4.0")
    print("✅ DocuGen模板识别和填充")
    print("✅ AI智能大纲和章节生成")
    print("✅ 简化RAG知识库功能")
    print("✅ 智谱GLM-4真实AI能力")
    print(f"📁 模板目录: {TEMPLATES_DIR}")
    print(f"📥 下载目录: {DOWNLOADS_DIR}")
    print(f"📚 知识库目录: {KNOWLEDGE_DIR}")
    print("🌐 监听地址: http://0.0.0.0:8000")
    print("📚 API文档: http://localhost:8000/docs")

    uvicorn.run(app, host="0.0.0.0", port=8000)
